# -*- coding: utf-8 -*-
"""
Llena la "Orden de Trabajo" de Mapfre y la convierte a PDF.

⚠️ REESCRITO con openpyxl (versión anterior usaba xlwt/xlutils, que NO
conservaba el logo de MAPFRE TEPEYAC incrustado en la plantilla — se
perdía al guardar). openpyxl SÍ conserva imágenes incrustadas al
leer/escribir .xlsx, confirmado con prueba real.

Requiere la plantilla en formato .xlsx (OT_MAPFRE_template.xlsx) — se
generó UNA VEZ convirtiendo el .xls original con LibreOffice:
    soffice --headless --convert-to xlsx OT_MAPFRE_template.xls

Requiere LibreOffice instalado en el servidor (soffice --headless) para
la conversión final xlsx -> pdf.
"""

import os
import subprocess
import tempfile
from datetime import date

import openpyxl

TEMPLATE_PATH = "OT_MAPFRE_template.xlsx"

# Datos fijos por grupo de pólizas — confirmados por Carlos.
# NOTA: razon_social sin salto de línea (\n) — se probó y el salto de
# línea hacía que solo se viera la segunda línea en el PDF, por la altura
# fija de la celda en la plantilla.
DATOS_FIJOS_POR_POLIZA = {
    "2612500003880": {
        "razon_social": "GRUPO ARRENDA MAX DE MORELOS S.A. de C.V.",
        "domicilio_col": "CHIPITLAN",
        "domicilio_estado": "MORELOS",
        "vigencia_desde": (1, 11, 2025),   # (DD, MM, AA)
        "vigencia_hasta": (1, 11, 2026),
        "instrucciones": "FAVOR DE DAR DE ALTA ASEGURADA",
    },
    "2612500003895": {
        "razon_social": "GRUPO ARRENDA MAX DE MORELOS S.A. de C.V.",
        "domicilio_col": "CHIPITLAN",
        "domicilio_estado": "MORELOS",
        "vigencia_desde": (1, 11, 2025),
        "vigencia_hasta": (1, 11, 2026),
        "instrucciones": "FAVOR DE DAR DE ALTA ASEGURADA",
    },
    "2612500003973": {
        "razon_social": "GRUPO ARRENDA MAX DE MORELOS S.A. de C.V.",
        "domicilio_col": "CHIPITLAN",
        "domicilio_estado": "MORELOS",
        "vigencia_desde": (1, 11, 2025),
        "vigencia_hasta": (1, 11, 2026),
        "instrucciones": "FAVOR DE DAR DE ALTA ASEGURADA",
    },
    "2612600000892": {  # CYSE/SNAC
        "razon_social": "GRUPO CONSULTOR Y ASESOR CYSE S.A. de C.V.",
        "domicilio_col": "CHIPITLAN",
        "domicilio_estado": "MORELOS",
        "vigencia_desde": (20, 2, 2026),
        "vigencia_hasta": (20, 2, 2027),
        "instrucciones": "FAVOR DE DAR DE ALTA  A LOS  ASEGURADOS.",
    },
}

AGENTE_CLAVE = "33910"
AGENTE_NOMBRE = "FRANCISCO JAVIER PORRAS VELAZQUEZ"

# Coordenadas confirmadas contra la plantilla real
CELDA_FECHA_SOLICITUD = "K1"
CELDA_NO_POLIZA = "K2"
CELDAS_VIGENCIA = ["A8", "B8", "C8", "D8", "E8", "F8"]  # DD MM AA DD MM AA
CELDA_RAZON_SOCIAL = "C12"
CELDA_DOMICILIO_COL = "K15"
CELDA_DOMICILIO_ESTADO = "G16"
CELDA_AGENTE_CLAVE = "C19"
CELDA_AGENTE_NOMBRE = "G19"
CELDA_INSTRUCCIONES = "A26"


def generar_orden_trabajo_mapfre(
    no_poliza: str,
    fecha_solicitud: date = None,
    template_path: str = TEMPLATE_PATH,
    salida_pdf: str = "/tmp/orden_trabajo_mapfre.pdf",
) -> str:
    datos = DATOS_FIJOS_POR_POLIZA.get(no_poliza)
    if not datos:
        raise ValueError(
            f"No hay datos fijos confirmados para la póliza {no_poliza!r}. "
            f"Pólizas soportadas: {list(DATOS_FIJOS_POR_POLIZA)}"
        )

    fecha_solicitud = fecha_solicitud or date.today()

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active  # "Hoja1"

    ws[CELDA_FECHA_SOLICITUD] = fecha_solicitud.strftime("%d/%m/%Y")
    ws[CELDA_NO_POLIZA] = no_poliza

    d1, m1, a1 = datos["vigencia_desde"]
    d2, m2, a2 = datos["vigencia_hasta"]
    for celda, valor in zip(CELDAS_VIGENCIA, [d1, m1, a1, d2, m2, a2]):
        ws[celda] = valor

    ws[CELDA_RAZON_SOCIAL] = datos["razon_social"]
    ws[CELDA_DOMICILIO_COL] = datos["domicilio_col"]
    ws[CELDA_DOMICILIO_ESTADO] = datos["domicilio_estado"]
    ws[CELDA_AGENTE_CLAVE] = AGENTE_CLAVE
    ws[CELDA_AGENTE_NOMBRE] = AGENTE_NOMBRE
    ws[CELDA_INSTRUCCIONES] = datos["instrucciones"]

    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_temp = os.path.join(tmpdir, "orden_trabajo.xlsx")
        wb.save(xlsx_temp)

        # Convertir a PDF con LibreOffice headless — SinglePageSheets
        # fuerza que quepa en una sola página.
        resultado = subprocess.run(
            [
                "soffice", "--headless", "--convert-to",
                'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}',
                "--outdir", tmpdir, xlsx_temp,
            ],
            capture_output=True, text=True, timeout=60,
        )
        if resultado.returncode != 0:
            raise RuntimeError(f"Error convirtiendo a PDF con LibreOffice: {resultado.stderr}")

        pdf_generado = os.path.join(tmpdir, "orden_trabajo.pdf")
        if not os.path.exists(pdf_generado):
            raise RuntimeError(
                f"LibreOffice no generó el PDF esperado. Salida: {resultado.stdout} {resultado.stderr}"
            )

        # La plantilla trae una segunda hoja ("ENDOSO", un ejemplo viejo
        # sin relación) que también se exporta — nos quedamos solo con la
        # página 1.
        from pypdf import PdfReader, PdfWriter
        lector = PdfReader(pdf_generado)
        escritor = PdfWriter()
        escritor.add_page(lector.pages[0])
        with open(salida_pdf, "wb") as f:
            escritor.write(f)

    return salida_pdf


if __name__ == "__main__":
    ruta = generar_orden_trabajo_mapfre(no_poliza="2612600000892", salida_pdf="/tmp/OT_MAPFRE_TEST.pdf")
    print("Generado:", ruta)
